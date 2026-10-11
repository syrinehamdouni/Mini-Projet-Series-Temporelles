"""Génère presentation/Presentation_KO_PEP.pptx.
- Le contenu méthodologique (ce que le prof exige) est toujours écrit.
- Les figures et tableaux réels sont insérés automatiquement s'ils existent (après exécution du pipeline) ;
  sinon un cadre 'à générer' apparaît. Relancer ce script après le pipeline pour obtenir la version complète."""
import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG, RES = ROOT / "figures", ROOT / "results"
KEY = json.loads((RES / "key_results.json").read_text()) if (RES / "key_results.json").exists() else {}
NAVY, GREY = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x44, 0x44, 0x44)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def txt(slide, x, y, w, h, lines, size=18, bold=False, color=GREY):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.font.size, p.font.bold, p.font.color.rgb = l, Pt(size), bold, color
        p.space_after = Pt(6)


def table(slide, csv, cols, x, y, w, nd=4, maxrows=9):
    p = RES / csv
    if not p.exists():
        txt(slide, x, y, w, 1, [f"[Tableau {csv} à générer : exécuter le pipeline]"], 14)
        return
    df = pd.read_csv(p)[cols].head(maxrows)
    t = slide.shapes.add_table(len(df) + 1, len(cols), Inches(x), Inches(y), Inches(w), Inches(0.35 * (len(df) + 1))).table
    for j, c in enumerate(cols):
        t.cell(0, j).text = c
    for i, row in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(row):
            t.cell(i, j).text = f"{v:.{nd}f}" if isinstance(v, float) else str(v)
    for r in t.rows:
        for c in r.cells:
            for p_ in c.text_frame.paragraphs:
                p_.font.size = Pt(11)


def add(title, bullets, fig=None, csv=None, cols=None, notes="", result=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    txt(s, 0.5, 0.3, 12.3, 0.9, [title], 30, True, NAVY)
    side = fig or csv
    lines = list(bullets) + ([f"► Résultat : {result}"] if result else [])
    txt(s, 0.5, 1.3, 6.3 if side else 12.3, 5.8, ["• " + b if not b.startswith("►") else b for b in lines], 16 if side else 18)
    if fig:
        if (FIG / fig).exists():
            s.shapes.add_picture(str(FIG / fig), Inches(7), Inches(1.4), width=Inches(6))
        else:
            txt(s, 7, 2.5, 6, 1, [f"[Figure {fig} à générer : exécuter le pipeline]"], 14)
    if csv:
        table(s, csv, cols, 7, 1.5, 6)
    s.notes_slide.notes_text_frame.text = notes
    return s


def k(name, fmt="{:.4f}"):
    v = KEY.get(name)
    return fmt.format(v) if isinstance(v, (int, float)) else (str(v) if v is not None else "n/d")


# ------------------------------------------------------------------------------------------------ slides
s = prs.slides.add_slide(prs.slide_layouts[6])
txt(s, 0.8, 2.2, 11.5, 1.5, ["Modélisation, prévision et diagnostic de séries temporelles multidimensionnelles"], 36, True, NAVY)
txt(s, 0.8, 4.0, 11.5, 1.5, ["Coca-Cola (KO) & PepsiCo (PEP) — données Yahoo Finance 2019-2024",
                              "Cours Séries Temporelles pour Data Science & IA — TEK-UP — Mini-projet", "Syrine"], 20)

add("Problématique et choix de la thématique",
    ["Thématique choisie : FINANCE (option A) — KO et PEP, deux actifs du même secteur, donc corrélés.",
     "Justification en une phrase : ce sont deux valeurs défensives très liquides et fortement liées, idéales pour tester VAR, cointégration et GARCH.",
     "Question : quel modèle prévoit le mieux le prix de KO à 10 jours, et les modèles sophistiqués battent-ils vraiment la baseline naïve ?",
     "Principe du barème : la rigueur du diagnostic prime sur le nombre de modèles."],
    notes="Rappeler que le choix de thématique est définitif et justifié en une phrase (consigne section 1).")

add("Données : source, période, transformations",
    ["Source : Yahoo Finance via yfinance ; journalier ; 2019-01-01 → 2024-12-31 (≈ 1 500 jours de bourse).",
     "Vecteur multidimensionnel : Open, High, Low, Close, Adj Close, Volume ; on modélise log(Adj Close).",
     "Adj Close = ajusté des dividendes et splits (transformation du fournisseur, documentée).",
     "Indicateurs dérivés : rendement log, moyenne mobile 20 j, volatilité réalisée 21 j.",
     "Split CHRONOLOGIQUE 90 % apprentissage / 10 % test — jamais aléatoire."],
    fig="fig01_prix.png", notes="Insister : split chronologique, graine 42, CSV en cache pour la reproductibilité.")

add("Analyse exploratoire : tendance, rendements, volatilité",
    ["Prix en niveau : tendance et rupture de mars 2020 (COVID).", "Rendements centrés sur 0 ; volatilité par grappes (clusters).",
     "Comparaison des échelles : prix bruts puis base 100.", "Statistiques descriptives : moyenne, variance, skewness, kurtosis (stats_descriptives.csv)."],
    fig="fig02_indicateurs.png", result=f"corrélation des rendements KO/PEP = {k('corr_rendements', '{:.2f}')}" if KEY else None)

add("ACF / PACF : brute puis différenciée",
    ["Série brute : ACF à décroissance lente → racine unitaire.", "Série différenciée : sert à identifier p (PACF) et q (ACF).",
     "L'interprétation est ÉCRITE dans le rapport, pas seulement le graphique."], fig="fig04_acf_pacf_diff.png",
    notes="Montrer aussi fig03 (niveau). Règle : PACF coupe → AR ; ACF coupe → MA.")

add("Décomposition STL et valeurs aberrantes",
    ["STL robuste (période 252) : tendance, saisonnalité, reste commentés.", "Aberrants : règle IQR ; chocs conservés ; indicateurs archivés pour audit, pas utilisés comme variables.",
     "Pas de suppression : le risque extrême est de l'information, GARCH est fait pour l'absorber."], fig="fig05_stl_KO.png",
    result=f"force de la tendance = {k('force_tendance', '{:.2f}')} ; force de la saisonnalité = {k('force_saisonnalite', '{:.2f}')}" if KEY else None)

add("Stationnarité : ADF ET KPSS",
    ["ADF : H0 = racine unitaire. KPSS : H0 = stationnaire → hypothèses opposées, donc complémentaires.",
     "Décision : ADF non rejeté + KPSS rejeté ⇒ non stationnaire ; l'inverse ⇒ stationnaire ; sinon ambigu.",
     "Conclusion explicite sur l'ordre de différenciation d."], csv="stationnarite.csv", cols=["serie", "ADF_p", "KPSS_p", "conclusion"],
    result=f"d = {k('d_ordre', '{:.0f}')}" if KEY else None)

add("ARIMA / SARIMA : identification et diagnostic",
    ["Ordres lus sur ACF/PACF PUIS confirmés par recherche automatique (pmdarima / auto.arima) : les deux approches sont comparées.",
     "SARIMA saisonnier hebdomadaire (m = 5) : amélioration réelle ou complexité inutile ?",
     "Diagnostic complet : Ljung-Box (indépendance), Jarque-Bera et Shapiro-Wilk (normalité), ARCH-LM (variance).",
     "Un modèle sans diagnostic n'est pas noté."], csv="diag_arima.csv", cols=["modele", "ordre", "AIC", "LB10_p", "JB_p", "ARCH10_p"],
    notes="Chaque p-value doit être interprétée : rejet ou non de H0, et conséquence pour le modèle.")

add("Diagnostic des résidus ARIMA", ["Résidus indépendants ? (ACF dans la bande, Ljung-Box).", "Normalité rejetée = queues épaisses (courant en finance).",
                                      "ARCH rejeté → variance non constante → justifie GARCH (section suivante)."], fig="fig08_resid_ARIMA_auto.png",
    result=f"meilleur modèle (AIC) = {k('best_arima')} {k('best_arima_order')} ; ARCH-LM p = {k('arima_arch_p')}" if KEY else None)

add("VAR et causalité de Granger",
    ["VAR sur les rendements (stationnaires), ordre choisi par BIC/AIC.", "Granger dans les deux sens : pouvoir prédictif, pas causalité économique.",
     "Diagnostic du VAR : stabilité, portemanteau, normalité, ARCH par équation ; expliciter tout échec de diagnostic."], fig="fig11_resid_VAR_KO.png",
    result=(f"VAR({k('var_order', '{:.0f}')}) ; Granger PEP→KO p = {k('granger_PEP_to_KO_p')} ; KO→PEP p = {k('granger_KO_to_PEP_p')}") if KEY else None)

add("Cointégration et VECM",
    ["Deux séries I(1) peuvent partager un équilibre de long terme.", "Comparer Engle-Granger et Johansen, y compris les divergences entre implémentations.",
     "Le VECM rank 1 est conditionné par Engle-Granger ; un rang Johansen plein signale une spécification à discuter.",
     "Le terme de correction d'erreur α décrit l'ajustement, pas une causalité économique."], fig="fig12_spread_cointegration.png",
    result=(f"Engle-Granger p = {k('engle_granger_p')} ; rang Johansen Python = {k('johansen_rank', '{:.0f}')}") if KEY else None)

add("ARCH-LM et GARCH",
    ["Test ARCH-LM sur les résidus du meilleur ARIMA : variance non constante ?", "GARCH(1,1) comparé à des ordres supérieurs et à la loi de Student (BIC).",
     "Résidus standardisés : doivent redevenir homoscédastiques.", "Persistance Σα+Σβ et demi-vie d'un choc de volatilité."],
    csv="garch_selection.csv", cols=["modele", "loi", "AIC", "BIC"],
    result=f"meilleur : {k('garch_best')} ; persistance = {k('garch_persistence')} ; demi-vie ≈ {k('garch_half_life_days', '{:.1f}')} j" if KEY else None)

add("Volatilité conditionnelle et événements", ["Pics reliés au krach COVID (mars 2020) et à la remontée des taux (2022).", "Retour progressif au calme = persistance.",
                                                 "Diagnostic des résidus standardisés (fig14)."], fig="fig13_garch_volatilite.png")

add("Réseaux récurrents : LSTM et GRU",
    ["Fenêtres glissantes (10 ou 30 jours) de [rendement KO, rendement PEP, volatilité 21 j] ; échelles réajustées dans chaque pli.",
     "Prévision directe des 10 prochains rendements ; 2 couches (32/16), dropout 0.1 ou 0.3, arrêt précoce.",
     "Hyperparamètres et époques choisis sur une validation chronologique imbriquée dans le train de chaque pli.",
     "Petit réseau car ~1 300 fenêtres : un gros réseau surapprendrait."], fig="fig15_dl_pertes.png", notes="Expliquer chaque hyperparamètre : fenêtre, couches, dropout.")

add("Validation croisée temporelle (walk-forward)",
    ["Jamais de k-fold classique sur une série temporelle.", "5 plis extensibles : entraînement sur le passé, test sur le bloc suivant, embargo de 10 jours.",
     "Comparaison à la baseline 'rendement nul' (marche aléatoire)."], fig="fig16_dl_cv_walkforward.png",
    result=f"RMSE moyen CV : {k('dl_cv_mean')} vs naïf {k('dl_cv_naive_mean')}" if KEY else None)

add("Prévision h = 10 et métriques (tableau unique)",
    ["Blocs de 10 jours consécutifs sur tout le test, mêmes observations pour tous les modèles.", "MSE, RMSE, MAE, MAPE + complexité, interprétabilité, temps de calcul.",
     "Baseline naïve incluse : sans elle, aucun gain ne peut être affirmé."], csv="metrics.csv", cols=["modele", "RMSE", "MAE", "MAPE", "complexite"],
    result=f"meilleur RMSE : {k('best_rmse_model')}" if KEY else None)

add("Test de Diebold-Mariano",
    ["Une différence de RMSE n'est pas forcément significative : on la teste.", "DM avec variance de Newey-West et correction de Harvey (perte quadratique).",
     "Lecture : p < 0.05 → précision significativement différente."], fig="fig19_dm_pvalues.png",
    result=f"significatifs vs naïf : {k('models_sig_vs_naive')}" if KEY else None)

add("Conclusion : meilleur modèle selon le cas d'usage",
    ["Critères : performance (RMSE + DM), interprétabilité, coût de calcul.", "Un modèle simple bien diagnostiqué vaut mieux qu'un modèle sophistiqué mal validé.",
     "Sur des actions liquides, battre la marche aléatoire est difficile : l'honnêteté sur ce point est valorisée."], fig="fig18_metriques.png")

add("IA générative : méthode (section 5)",
    ["1. Hypothèses de modèles générées à partir des statistiques descriptives (prompt + réponse documentés intégralement).",
     "2. Vulgarisation pour un décideur non-statisticien (annexe).", "3. Recommandations simulées avec limites formulées par le LLM puis complétées.",
     "Les analyses humaines et réponses existantes sont préservées ; l'ordre chronologique réel des fichiers ne peut pas être certifié.",
     "Application Streamlit : prompts, réponses et comparaison critique."])

add("IA générative : esprit critique",
    ["Hallucinations statistiques : p-values inventées, causalités non vérifiées, simplifications abusives.", "Chaque affirmation chiffrée du LLM est vérifiée contre results/*.csv.",
     "Tableau : affirmation → vérification → verdict → correction.", "Bilan : temps/clarté gagnés vs ce que le LLM n'a pas su faire."])

add("Limites et discussion critique",
    ["Hypothèses violées : normalité des résidus, variance non constante (traitée par GARCH).", "Un seul actif prévu, 6 ans de données, un seul régime de marché dominant.",
     "Estimations non stables dans le temps ; DM approximé sur des blocs de 10 jours ; réseaux non ré-entraînés à chaque bloc.",
     "Prévoir un prix est quasi impossible (efficience) : les recommandations simulées ne sont pas un conseil financier.",
     "IA générative : gain de temps réel, mais risque de confiance excessive sans vérification."])

add("Livrables et reproductibilité",
    ["Code Python (socle classique + avancé) ET R (socle classique), commentés : quoi + pourquoi.", "Notebook Jupyter exécutable de bout en bout (graine 42, versions affichées).",
     "Rapport 15-25 pages, application Streamlit, fiche syrine.txt.", "Remise : 12 octobre 2026, 23h55, par mail à ahmed.dhouibi@tek-up.de (archive .zip)."])

add("Couverture du barème (20 points)",
    ["A — Fondamentaux (10) : données, exploration, ARIMA/SARIMA/VAR avec diagnostics, évaluation.", "B — Modélisation avancée (4) : GARCH diagnostiqué, LSTM/GRU + CV temporelle.",
     "C — IA générative (4) : hypothèses, vulgarisation, recommandations avec limites, comparaison humain/LLM.", "D — Approfondissement (2) : cointégration/VECM."])

out = ROOT / "presentation" / "Presentation_KO_PEP.pptx"
prs.save(out)
print(f"{len(prs.slides)} diapositives -> {out}")
